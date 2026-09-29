from __future__ import annotations

import base64
import io
import os
import re
from pathlib import Path
from typing import Optional

from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.lib import colors

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


# ============================================================
# FONT CONFIGURATION
# ============================================================

_registered_fonts: set[str] = set()


def find_unicode_font() -> Optional[str]:
    """
    Find a Windows font that can be used by ReportLab.

    We intentionally do not depend on DejaVuSans because it may
    not be installed on every Windows machine.
    """

    possible_fonts = [
        # Arial
        r"C:\Windows\Fonts\arial.ttf",

        # Arial Unicode, if installed
        r"C:\Windows\Fonts\ARIALUNI.TTF",

        # Calibri
        r"C:\Windows\Fonts\calibri.ttf",

        # Segoe UI
        r"C:\Windows\Fonts\segoeui.ttf",

        # Nirmala UI - useful for Indian scripts
        r"C:\Windows\Fonts\Nirmala.ttf",
        r"C:\Windows\Fonts\NirmalaUI.ttf",

        # Cambria
        r"C:\Windows\Fonts\cambria.ttc",

        # Tahoma
        r"C:\Windows\Fonts\tahoma.ttf",
    ]

    for font_path in possible_fonts:
        if os.path.exists(font_path):
            return font_path

    return None


def find_unicode_bold_font() -> Optional[str]:
    """
    Find a bold Unicode font available on Windows.
    """

    possible_fonts = [
        r"C:\Windows\Fonts\arialbd.ttf",
        r"C:\Windows\Fonts\ARIALBD.TTF",
        r"C:\Windows\Fonts\calibrib.ttf",
        r"C:\Windows\Fonts\segoeuib.ttf",
        r"C:\Windows\Fonts\NirmalaB.ttf",
        r"C:\Windows\Fonts\NirmalaUI-Bold.ttf",
        r"C:\Windows\Fonts\tahomabd.ttf",
    ]

    for font_path in possible_fonts:
        if os.path.exists(font_path):
            return font_path

    return None


def register_pdf_fonts() -> tuple[str, str]:
    """
    Register Unicode fonts for ReportLab.

    Returns:
        (regular_font_name, bold_font_name)
    """

    regular_path = find_unicode_font()
    bold_path = find_unicode_bold_font()

    regular_name = "LegalEaseUnicode"
    bold_name = "LegalEaseUnicodeBold"

    # Regular font
    if regular_path:

        if regular_name not in _registered_fonts:

            try:
                pdfmetrics.registerFont(
                    TTFont(
                        regular_name,
                        regular_path
                    )
                )

                _registered_fonts.add(
                    regular_name
                )

            except Exception:
                regular_name = "Helvetica"

    else:
        regular_name = "Helvetica"

    # Bold font
    if bold_path:

        if bold_name not in _registered_fonts:

            try:
                pdfmetrics.registerFont(
                    TTFont(
                        bold_name,
                        bold_path
                    )
                )

                _registered_fonts.add(
                    bold_name
                )

            except Exception:
                bold_name = regular_name

    else:
        bold_name = regular_name

    return regular_name, bold_name


# Register fonts when this module loads.
PDF_FONT, PDF_BOLD_FONT = register_pdf_fonts()


# ============================================================
# GENERAL TEXT HELPERS
# ============================================================

def clean_text(text: str) -> str:
    """
    Clean text without destroying useful Unicode characters.
    """

    if text is None:
        return ""

    text = str(text)

    # Normalize unusual line endings.
    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    return text.strip()


def escape_pdf_text(text: str) -> str:
    """
    Escape text before putting it inside a ReportLab Paragraph.
    """

    text = clean_text(text)

    # ReportLab Paragraph interprets XML-like characters.
    text = (
        text
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )

    return text


def markdown_to_pdf_text(text: str) -> str:
    """
    Convert a small subset of Markdown formatting to
    ReportLab-compatible markup.
    """

    text = escape_pdf_text(text)

    # Bold: **text**
    text = re.sub(
        r"\*\*(.+?)\*\*",
        r"<b>\1</b>",
        text
    )

    # Italic: *text*
    text = re.sub(
        r"(?<!\*)\*([^*]+)\*(?!\*)",
        r"<i>\1</i>",
        text
    )

    return text


# ============================================================
# TXT EXPORT
# ============================================================

def format_txt(text: str) -> bytes:
    """
    Return the document as UTF-8 text.

    UTF-8 ensures characters such as ₹ are preserved.
    """

    text = clean_text(text)

    return text.encode(
        "utf-8"
    )


# ============================================================
# DOCX EXPORT
# ============================================================

def add_logo_to_docx(
    document: Document,
    logo_base64: Optional[str]
) -> None:

    if not logo_base64:
        return

    try:

        # Remove possible data URL prefix.
        if "," in logo_base64:
            logo_base64 = logo_base64.split(
                ",",
                1
            )[1]

        logo_bytes = base64.b64decode(
            logo_base64
        )

        image_stream = io.BytesIO(
            logo_bytes
        )

        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = paragraph.add_run()

        run.add_picture(
            image_stream,
            width=Inches(1.2)
        )

    except Exception:
        # Logo should never prevent document generation.
        pass


def add_docx_line(
    document: Document,
    line: str
) -> None:

    line = line.strip()

    if not line:
        document.add_paragraph()
        return

    # Remove Markdown bold markers for DOCX.
    clean_line = line.replace(
        "**",
        ""
    )

    # Heading detection.
    if line.startswith("# "):

        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = paragraph.add_run(
            clean_line[2:].strip()
        )

        run.bold = True
        run.font.size = Pt(18)

        return

    if line.startswith("## "):

        paragraph = document.add_paragraph()

        run = paragraph.add_run(
            clean_line[3:].strip()
        )

        run.bold = True
        run.font.size = Pt(15)

        return

    if line.startswith("### "):

        paragraph = document.add_paragraph()

        run = paragraph.add_run(
            clean_line[4:].strip()
        )

        run.bold = True
        run.font.size = Pt(13)

        return

    # Horizontal rule.
    if line in {
        "---",
        "***",
        "___"
    }:

        paragraph = document.add_paragraph(
            "____________________________________________________________"
        )

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        return

    # Normal paragraph.
    paragraph = document.add_paragraph()

    run = paragraph.add_run(
        clean_line
    )

    run.font.name = "Arial"
    run.font.size = Pt(11)


def format_docx(
    text: str,
    document_type: str = "Legal Document",
    brand_name: str = "LegalEase",
    logo_base64: Optional[str] = None
) -> bytes:

    document = Document()

    # --------------------------------------------------------
    # Page margins
    # --------------------------------------------------------

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    # --------------------------------------------------------
    # Logo
    # --------------------------------------------------------

    add_logo_to_docx(
        document,
        logo_base64
    )

    # --------------------------------------------------------
    # Brand name
    # --------------------------------------------------------

    if brand_name:

        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = paragraph.add_run(
            brand_name
        )

        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(10)

    # --------------------------------------------------------
    # Document content
    # --------------------------------------------------------

    text = clean_text(text)

    lines = text.split("\n")

    for line in lines:

        add_docx_line(
            document,
            line
        )

    # --------------------------------------------------------
    # Save to memory
    # --------------------------------------------------------

    output = io.BytesIO()

    document.save(
        output
    )

    return output.getvalue()


# ============================================================
# PDF HELPERS
# ============================================================

def create_pdf_styles():

    body_style = ParagraphStyle(
        name="LegalEaseBody",

        fontName=PDF_FONT,

        fontSize=10.5,

        leading=15,

        alignment=TA_LEFT,

        spaceAfter=7,

        wordWrap="LTR",
    )

    heading_style = ParagraphStyle(
        name="LegalEaseHeading",

        fontName=PDF_BOLD_FONT,

        fontSize=13,

        leading=17,

        alignment=TA_LEFT,

        spaceBefore=10,

        spaceAfter=7,

        wordWrap="LTR",
    )

    title_style = ParagraphStyle(
        name="LegalEaseTitle",

        fontName=PDF_BOLD_FONT,

        fontSize=17,

        leading=22,

        alignment=TA_CENTER,

        spaceAfter=15,

        wordWrap="LTR",
    )

    subtitle_style = ParagraphStyle(
        name="LegalEaseSubtitle",

        fontName=PDF_FONT,

        fontSize=9,

        leading=12,

        alignment=TA_CENTER,

        spaceAfter=12,

        wordWrap="LTR",
    )

    return (
        body_style,
        heading_style,
        title_style,
        subtitle_style,
    )


def make_pdf_paragraph(
    line: str,
    body_style: ParagraphStyle,
    heading_style: ParagraphStyle,
    title_style: ParagraphStyle
):

    stripped = line.strip()

    if not stripped:

        return Spacer(
            1,
            5
        )

    # Markdown H1
    if stripped.startswith("# "):

        content = stripped[2:].strip()

        return Paragraph(
            markdown_to_pdf_text(content),
            title_style
        )

    # Markdown H2
    if stripped.startswith("## "):

        content = stripped[3:].strip()

        return Paragraph(
            markdown_to_pdf_text(content),
            heading_style
        )

    # Markdown H3
    if stripped.startswith("### "):

        content = stripped[4:].strip()

        return Paragraph(
            markdown_to_pdf_text(content),
            heading_style
        )

    # Numbered section headings.
    if re.match(
        r"^\d+[\.\)]\s+[A-Z]",
        stripped
    ):

        return Paragraph(
            markdown_to_pdf_text(stripped),
            heading_style
        )

    # ALL CAPS short headings.
    if (
        len(stripped) < 100
        and stripped.upper() == stripped
        and any(char.isalpha() for char in stripped)
    ):

        return Paragraph(
            markdown_to_pdf_text(stripped),
            heading_style
        )

    # Signature lines and ordinary content.
    return Paragraph(
        markdown_to_pdf_text(stripped),
        body_style
    )


# ============================================================
# PDF EXPORT
# ============================================================

def format_pdf(
    text: str,
    document_type: str = "Legal Document",
    brand_name: str = "LegalEase",
    logo_base64: Optional[str] = None
) -> bytes:

    """
    Generate a Unicode-capable PDF.

    Important:
    We use a real TTF font when available instead of
    Helvetica so characters such as ₹ can be rendered.
    """

    text = clean_text(text)

    output = io.BytesIO()

    document = SimpleDocTemplate(

        output,

        pagesize=A4,

        rightMargin=20 * mm,

        leftMargin=20 * mm,

        topMargin=20 * mm,

        bottomMargin=20 * mm,

        title=document_type,

        author=brand_name or "LegalEase",

    )

    (
        body_style,
        heading_style,
        title_style,
        subtitle_style,
    ) = create_pdf_styles()

    story = []

    # --------------------------------------------------------
    # Optional logo
    # --------------------------------------------------------

    if logo_base64:

        try:

            from reportlab.platypus import Image

            if "," in logo_base64:
                logo_base64 = (
                    logo_base64.split(
                        ",",
                        1
                    )[1]
                )

            logo_bytes = base64.b64decode(
                logo_base64
            )

            logo_stream = io.BytesIO(
                logo_bytes
            )

            image = Image(
                logo_stream,
                width=30 * mm,
                height=30 * mm
            )

            image.hAlign = "CENTER"

            story.append(image)

            story.append(
                Spacer(
                    1,
                    5
                )
            )

        except Exception:
            # Ignore invalid logos.
            pass

    # --------------------------------------------------------
    # Brand name
    # --------------------------------------------------------

    if brand_name:

        story.append(
            Paragraph(
                escape_pdf_text(
                    brand_name
                ),
                subtitle_style
            )
        )

    # --------------------------------------------------------
    # Document content
    # --------------------------------------------------------

    lines = text.split("\n")

    for line in lines:

        element = make_pdf_paragraph(

            line,

            body_style,

            heading_style,

            title_style
        )

        story.append(
            element
        )

    # --------------------------------------------------------
    # Footer
    # --------------------------------------------------------

    def add_page_number(canvas, doc):

        canvas.saveState()

        canvas.setFont(
            PDF_FONT,
            8
        )

        page_number = canvas.getPageNumber()

        footer = (
            f"{brand_name or 'LegalEase'} | "
            f"Page {page_number}"
        )

        canvas.drawCentredString(
            A4[0] / 2,
            10 * mm,
            footer
        )

        canvas.restoreState()

    # --------------------------------------------------------
    # Build PDF
    # --------------------------------------------------------

    document.build(
        story,
        onFirstPage=add_page_number,
        onLaterPages=add_page_number
    )

    return output.getvalue()